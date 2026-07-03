"""
Excel导出模块
将匹配结果导出为格式化的Excel文件，支持排序和统计
"""

import os
from typing import Dict, List, Optional
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

from utils import setup_logger, format_duration

logger = setup_logger("excel_exporter")


def export_to_excel(
    match_result: Dict[str, Dict],
    output_path: str,
    pdf_name: str = "",
    stats: Optional[Dict] = None,
    duration: float = 0
) -> str:
    """
    导出匹配结果到Excel文件
    
    Args:
        match_result: 匹配结果字典，键为人名，值为包含 count 和 pages 的字典
        output_path: 输出文件路径
        pdf_name: PDF文件名（用于统计信息）
        stats: 统计信息字典
        duration: 匹配耗时（秒）
    
    Returns:
        输出文件的完整路径
    
    Raises:
        ValueError: 结果为空
    """
    if not match_result:
        raise ValueError("匹配结果为空")
    
    # 创建工作簿
    wb = Workbook()
    
    # 创建匹配结果工作表
    ws_results = wb.active
    ws_results.title = "匹配结果"
    
    # 按出现次数降序排序
    sorted_results = sorted(
        match_result.items(),
        key=lambda x: x[1]['count'],
        reverse=True
    )
    
    # 写入表头
    headers = ["名称", "出现次数", "出现页码"]
    header_font = Font(bold=True, size=12)
    header_fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    
    for col, header in enumerate(headers, 1):
        cell = ws_results.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
    
    # 写入数据
    data_font = Font(size=11)
    data_alignment = Alignment(vertical="center")
    pages_alignment = Alignment(vertical="center", wrap_text=True)
    
    for row_idx, (name, data) in enumerate(sorted_results, 2):
        # 名称
        cell_name = ws_results.cell(row=row_idx, column=1, value=name)
        cell_name.font = data_font
        cell_name.alignment = data_alignment
        
        # 出现次数
        cell_count = ws_results.cell(row=row_idx, column=2, value=data['count'])
        cell_count.font = data_font
        cell_count.alignment = Alignment(horizontal="center", vertical="center")
        
        # 出现页码
        pages_str = ", ".join(str(p) for p in data['pages'])
        cell_pages = ws_results.cell(row=row_idx, column=3, value=pages_str)
        cell_pages.font = data_font
        cell_pages.alignment = pages_alignment
    
    # 设置列宽
    ws_results.column_dimensions['A'].width = 15
    ws_results.column_dimensions['B'].width = 12
    ws_results.column_dimensions['C'].width = 50
    
    # 添加边框
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    for row in ws_results.iter_rows(min_row=1, max_row=len(sorted_results) + 1, max_col=3):
        for cell in row:
            cell.border = thin_border
    
    # 创建统计信息工作表
    if stats:
        create_summary_sheet(wb, stats, pdf_name, duration)
    
    # 保存文件
    wb.save(output_path)
    logger.info(f"已导出Excel: {output_path}")
    
    return output_path


def create_summary_sheet(
    wb: Workbook,
    stats: Dict,
    pdf_name: str = "",
    duration: float = 0
) -> None:
    """
    创建统计信息工作表
    
    Args:
        wb: 工作簿对象
        stats: 统计信息字典
        pdf_name: PDF文件名
        duration: 匹配耗时（秒）
    """
    ws_summary = wb.create_sheet(title="统计信息", index=0)
    
    # 标题样式
    title_font = Font(bold=True, size=14)
    label_font = Font(bold=True, size=11)
    value_font = Font(size=11)
    
    # 写入标题
    ws_summary.cell(row=1, column=1, value="PDF人名匹配 - 统计报告").font = title_font
    ws_summary.merge_cells('A1:B1')
    
    # 写入生成时间
    ws_summary.cell(row=2, column=1, value="生成时间:").font = label_font
    ws_summary.cell(row=2, column=2, value=datetime.now().strftime("%Y-%m-%d %H:%M:%S")).font = value_font
    
    # 写入PDF信息
    row = 4
    ws_summary.cell(row=row, column=1, value="PDF文件:").font = label_font
    ws_summary.cell(row=row, column=2, value=pdf_name).font = value_font
    
    row += 1
    ws_summary.cell(row=row, column=1, value="匹配耗时:").font = label_font
    ws_summary.cell(row=row, column=2, value=format_duration(duration)).font = value_font
    
    # 写入统计信息
    row += 2
    ws_summary.cell(row=row, column=1, value="匹配统计").font = Font(bold=True, size=12)
    row += 1
    
    stats_items = [
        ("人名库总数", stats.get("total_names", 0)),
        ("匹配到的人名数", stats.get("matched_count", 0)),
        ("未匹配到的人名数", stats.get("unmatched_count", 0)),
        ("总匹配次数", stats.get("total_matches", 0)),
    ]
    
    for label, value in stats_items:
        ws_summary.cell(row=row, column=1, value=label).font = label_font
        ws_summary.cell(row=row, column=2, value=value).font = value_font
        row += 1
    
    # 写入匹配到的人名
    row += 1
    ws_summary.cell(row=row, column=1, value="匹配到的人名:").font = label_font
    row += 1
    
    matched_names = stats.get("matched_names", [])
    if matched_names:
        # 每行显示5个名字
        names_per_row = 5
        for i in range(0, len(matched_names), names_per_row):
            batch = matched_names[i:i + names_per_row]
            ws_summary.cell(row=row, column=1, value=", ".join(batch)).font = value_font
            row += 1
    
    # 写入未匹配到的人名
    row += 1
    ws_summary.cell(row=row, column=1, value="未匹配到的人名:").font = label_font
    row += 1
    
    unmatched_names = stats.get("unmatched_names", [])
    if unmatched_names:
        names_per_row = 5
        for i in range(0, len(unmatched_names), names_per_row):
            batch = unmatched_names[i:i + names_per_row]
            ws_summary.cell(row=row, column=1, value=", ".join(batch)).font = value_font
            row += 1
    
    # 设置列宽
    ws_summary.column_dimensions['A'].width = 20
    ws_summary.column_dimensions['B'].width = 40


def get_output_filename(pdf_path: str) -> str:
    """
    根据PDF文件名生成输出Excel文件名
    
    Args:
        pdf_path: PDF文件路径
    
    Returns:
        输出文件路径
    """
    # 获取PDF文件名（不含扩展名）
    pdf_name = os.path.splitext(os.path.basename(pdf_path))[0]
    
    # 生成时间戳
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    
    # 生成输出文件名
    output_name = f"{pdf_name}_人名匹配结果_{timestamp}.xlsx"
    
    # 获取PDF所在目录
    pdf_dir = os.path.dirname(pdf_path)
    
    return os.path.join(pdf_dir, output_name)
