"""
多格式导出模块 v2.0
支持 Excel/CSV/JSON 三种格式
"""

import os
import json
import csv
from typing import Dict, List, Optional
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

from stats import AnalysisStats, StatsCalculator
from context_extractor import ContextExtractor
from utils import setup_logger

logger = setup_logger("exporter")


class Exporter:
    """多格式导出器"""
    
    @staticmethod
    def export(
        match_result,
        output_path: str,
        formats: Dict[str, bool] = None,
        stats: Optional[AnalysisStats] = None,
        include_context: bool = True,
        context_length: int = 30,
        sort_by: str = "count",
        sort_order: str = "desc"
    ) -> List[str]:
        """
        导出匹配结果
        
        Args:
            match_result: MatchResult对象
            output_path: 输出路径（基础文件名，不含扩展名）
            formats: 导出格式 {'excel': True, 'csv': True, 'json': False}
            stats: 统计信息对象
            include_context: 是否包含上下文
            context_length: 上下文长度
            sort_by: 排序字段
            sort_order: 排序方式
        
        Returns:
            实际生成的文件路径列表
        """
        if formats is None:
            formats = {'excel': True, 'csv': False, 'json': False}
        
        # 获取表格数据
        table_data = StatsCalculator.get_keyword_table(match_result, sort_by, sort_order)
        
        # 获取上下文数据
        context_data = None
        if include_context:
            context_data = ContextExtractor.build_context_table(match_result, context_length)
        
        generated_files = []
        
        # Excel导出
        if formats.get('excel'):
            excel_path = output_path + '.xlsx'
            Exporter._export_excel(table_data, context_data, stats, excel_path)
            generated_files.append(excel_path)
        
        # CSV导出
        if formats.get('csv'):
            csv_path = output_path + '.csv'
            Exporter._export_csv(table_data, csv_path)
            generated_files.append(csv_path)
            
            # 上下文单独导出为CSV
            if context_data:
                ctx_csv_path = output_path + '_上下文.csv'
                Exporter._export_context_csv(context_data, ctx_csv_path)
                generated_files.append(ctx_csv_path)
        
        # JSON导出
        if formats.get('json'):
            json_path = output_path + '.json'
            Exporter._export_json(table_data, context_data, stats, match_result, json_path)
            generated_files.append(json_path)
        
        logger.info(f"导出完成: {', '.join(os.path.basename(f) for f in generated_files)}")
        return generated_files
    
    @staticmethod
    def _export_excel(
        table_data: List[dict],
        context_data: Optional[List[dict]],
        stats: Optional[AnalysisStats],
        output_path: str
    ) -> None:
        """导出Excel格式"""
        wb = Workbook()
        
        # 创建关键词结果工作表
        ws_result = wb.active
        ws_result.title = "关键词匹配结果"
        Exporter._write_result_sheet(ws_result, table_data)
        
        # 创建上下文工作表
        if context_data:
            ws_context = wb.create_sheet("上下文详情")
            Exporter._write_context_sheet(ws_context, context_data)
        
        # 创建统计工作表
        if stats:
            ws_stats = wb.create_sheet("统计信息")
            Exporter._write_stats_sheet(ws_stats, stats)
        
        # 保存
        wb.save(output_path)
        logger.info(f"Excel导出: {output_path}")
    
    @staticmethod
    def _write_result_sheet(ws, table_data: List[dict]) -> None:
        """写入结果工作表"""
        # 表头样式
        header_font = Font(bold=True, size=11)
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font_white = Font(bold=True, size=11, color="FFFFFF")
        header_alignment = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # 写入表头
        headers = ["关键词", "出现次数", "出现页码", "首次出现", "最后出现", "PDF数量", "PDF文件"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font_white
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # 写入数据
        for row_idx, data in enumerate(table_data, 2):
            row_data = [
                data['keyword'],
                data['count'],
                data['pages'],
                data['first_page'] if data['first_page'] > 0 else '',
                data['last_page'] if data['last_page'] > 0 else '',
                data['pdf_count'],
                data['pdfs']
            ]
            
            for col, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col, value=value)
                cell.border = thin_border
                if col in [2, 4, 5, 6]:  # 数字列居中
                    cell.alignment = Alignment(horizontal="center")
        
        # 设置列宽
        ws.column_dimensions['A'].width = 15
        ws.column_dimensions['B'].width = 12
        ws.column_dimensions['C'].width = 30
        ws.column_dimensions['D'].width = 12
        ws.column_dimensions['E'].width = 12
        ws.column_dimensions['F'].width = 12
        ws.column_dimensions['G'].width = 30
    
    @staticmethod
    def _write_context_sheet(ws, context_data: List[dict]) -> None:
        """写入上下文工作表"""
        header_font = Font(bold=True, size=11)
        header_fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
        header_font_white = Font(bold=True, size=11, color="FFFFFF")
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        headers = ["关键词", "PDF文件", "页码", "上下文"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font_white
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border
        
        for row_idx, data in enumerate(context_data, 2):
            values = [data['keyword'], data['pdf_name'], data['page_num'], data['context']]
            for col, value in enumerate(values, 1):
                cell = ws.cell(row=row_idx, column=col, value=value)
                cell.border = thin_border
                if col == 3:
                    cell.alignment = Alignment(horizontal="center")
        
        ws.column_dimensions['A'].width = 15
        ws.column_dimensions['B'].width = 25
        ws.column_dimensions['C'].width = 10
        ws.column_dimensions['D'].width = 80
    
    @staticmethod
    def _write_stats_sheet(ws, stats: AnalysisStats) -> None:
        """写入统计工作表"""
        header_font = Font(bold=True, size=12)
        label_font = Font(bold=True, size=11)
        value_font = Font(size=11)
        
        row = 1
        ws.cell(row=row, column=1, value="PDF关键词分析 - 统计报告").font = Font(bold=True, size=14)
        ws.merge_cells('A1:B1')
        
        row = 2
        ws.cell(row=row, column=1, value=f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}").font = Font(size=10, color="666666")
        
        # PDF统计
        row = 4
        ws.cell(row=row, column=1, value="PDF统计").font = header_font
        
        pdf_stats = [
            ("PDF数量", stats.pdf.total_pdfs),
            ("总页数", stats.pdf.total_pages),
            ("有文字页数", stats.pdf.pages_with_text),
        ]
        if stats.pdf.pages_needing_ocr > 0:
            pdf_stats.append(("需要OCR页数", stats.pdf.pages_needing_ocr))
        
        for label, value in pdf_stats:
            row += 1
            ws.cell(row=row, column=1, value=label).font = label_font
            ws.cell(row=row, column=2, value=value).font = value_font
        
        # 关键词统计
        row += 2
        ws.cell(row=row, column=1, value="关键词统计").font = header_font
        
        kw_stats = [
            ("关键词总数", stats.keyword.total_keywords),
            ("匹配到", stats.keyword.matched_keywords),
            ("未匹配", stats.keyword.unmatched_keywords),
            ("总命中次数", stats.keyword.total_hits),
        ]
        if stats.keyword.total_keywords > 0:
            kw_stats.append(("平均命中", f"{stats.keyword.avg_hits_per_keyword:.1f} 次/关键词"))
        
        for label, value in kw_stats:
            row += 1
            ws.cell(row=row, column=1, value=label).font = label_font
            ws.cell(row=row, column=2, value=value).font = value_font
        
        # 性能统计
        row += 2
        ws.cell(row=row, column=1, value="性能统计").font = header_font
        
        from utils import format_duration
        perf_stats = [
            ("总耗时", format_duration(stats.performance.total_time)),
        ]
        if stats.performance.pages_per_second > 0:
            perf_stats.append(("处理速度", f"{stats.performance.pages_per_second:.1f} 页/秒"))
        
        for label, value in perf_stats:
            row += 1
            ws.cell(row=row, column=1, value=label).font = label_font
            ws.cell(row=row, column=2, value=value).font = value_font
        
        # 命中最多的关键词
        if stats.keyword.top_keywords:
            row += 2
            ws.cell(row=row, column=1, value="命中最多的关键词").font = header_font
            
            for i, kw in enumerate(stats.keyword.top_keywords[:10], 1):
                row += 1
                ws.cell(row=row, column=1, value=f"{i}. {kw['name']}").font = value_font
                ws.cell(row=row, column=2, value=f"{kw['count']} 次").font = value_font
        
        # 未匹配的关键词
        if stats.keyword.unmatched_list:
            row += 2
            ws.cell(row=row, column=1, value="未匹配的关键词").font = header_font
            row += 1
            ws.cell(row=row, column=1, value=', '.join(stats.keyword.unmatched_list[:20])).font = value_font
        
        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 40
    
    @staticmethod
    def _export_csv(table_data: List[dict], output_path: str) -> None:
        """导出CSV格式"""
        with open(output_path, 'w', encoding='utf-8-sig', newline='') as f:
            writer = csv.writer(f)
            writer.writerow(["关键词", "出现次数", "出现页码", "首次出现", "最后出现", "PDF数量", "PDF文件"])
            
            for data in table_data:
                writer.writerow([
                    data['keyword'],
                    data['count'],
                    data['pages'],
                    data['first_page'] if data['first_page'] > 0 else '',
                    data['last_page'] if data['last_page'] > 0 else '',
                    data['pdf_count'],
                    data['pdfs']
                ])
        
        logger.info(f"CSV导出: {output_path}")
    
    @staticmethod
    def _export_context_csv(context_data: List[dict], output_path: str) -> None:
        """导出上下文CSV"""
        with open(output_path, 'w', encoding='utf-8-sig', newline='') as f:
            writer = csv.writer(f)
            writer.writerow(["关键词", "PDF文件", "页码", "上下文"])
            
            for data in context_data:
                writer.writerow([
                    data['keyword'],
                    data['pdf_name'],
                    data['page_num'],
                    data['context']
                ])
        
        logger.info(f"上下文CSV导出: {output_path}")
    
    @staticmethod
    def _export_json(
        table_data: List[dict],
        context_data: Optional[List[dict]],
        stats: Optional[AnalysisStats],
        match_result,
        output_path: str
    ) -> None:
        """导出JSON格式"""
        export_data = {
            'generated_at': datetime.now().isoformat(),
            'keywords': table_data,
        }
        
        if context_data:
            export_data['contexts'] = context_data
        
        if stats:
            export_data['statistics'] = stats.to_dict()
        
        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump(export_data, f, ensure_ascii=False, indent=2)
        
        logger.info(f"JSON导出: {output_path}")
    
    @staticmethod
    def get_output_filename(pdf_input: str, suffix: str = "_分析结果") -> str:
        """
        生成输出文件名
        
        Args:
            pdf_input: PDF输入路径
            suffix: 文件名后缀
        
        Returns:
            输出文件路径（不含扩展名）
        """
        if os.path.isdir(pdf_input):
            # 文件夹：使用文件夹名
            base_name = os.path.basename(pdf_input)
        elif os.path.isfile(pdf_input):
            # 单个文件：使用文件名
            base_name = os.path.splitext(os.path.basename(pdf_input))[0]
        else:
            # 多个文件：使用"批量分析"
            base_name = "批量分析"
        
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_dir = os.path.dirname(pdf_input) if os.path.isfile(pdf_input) else pdf_input
        
        return os.path.join(output_dir, f"{base_name}{suffix}_{timestamp}")
