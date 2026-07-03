"""
基准测试脚本
测试所有模块的功能和边界情况
"""

import sys
import os
import shutil

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

print('=' * 60)
print('基准测试 - 开始')
print('=' * 60)

errors = []
test_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), '_bench_temp')


def test(name, func):
    try:
        func()
        print(f'  ✓ {name}')
    except Exception as e:
        print(f'  ✗ {name}: {e}')
        errors.append((name, str(e)))


# 准备测试目录
if os.path.exists(test_dir):
    shutil.rmtree(test_dir)
os.makedirs(test_dir, exist_ok=True)

# ==================== utils.py ====================
print('\n[1/5] utils.py')

from utils import read_name_list, normalize_text, format_duration, generate_cache_path


def test_read_comma():
    f = os.path.join(test_dir, 'c.txt')
    with open(f, 'w', encoding='utf-8') as fp:
        fp.write('A,B,C')
    assert read_name_list(f) == ['A', 'B', 'C']


def test_read_newline():
    f = os.path.join(test_dir, 'n.txt')
    with open(f, 'w', encoding='utf-8') as fp:
        fp.write('A\nB\nC')
    assert read_name_list(f) == ['A', 'B', 'C']


def test_read_mixed():
    f = os.path.join(test_dir, 'm.txt')
    with open(f, 'w', encoding='utf-8') as fp:
        fp.write('A,B;C\nD')
    result = read_name_list(f)
    assert result == ['A', 'B', 'C', 'D'], f'实际: {result}'


def test_read_dup():
    f = os.path.join(test_dir, 'd.txt')
    with open(f, 'w', encoding='utf-8') as fp:
        fp.write('A,B,A,C,B')
    assert read_name_list(f) == ['A', 'B', 'C']


def test_read_semicolon():
    f = os.path.join(test_dir, 's.txt')
    with open(f, 'w', encoding='utf-8') as fp:
        fp.write('A;B;C')
    assert read_name_list(f) == ['A', 'B', 'C']


def test_read_chinese_comma():
    f = os.path.join(test_dir, 'cc.txt')
    with open(f, 'w', encoding='utf-8') as fp:
        fp.write('A，B，C')
    assert read_name_list(f) == ['A', 'B', 'C']


def test_read_chinese_semicolon():
    f = os.path.join(test_dir, 'cs.txt')
    with open(f, 'w', encoding='utf-8') as fp:
        fp.write('A；B；C')
    assert read_name_list(f) == ['A', 'B', 'C']


def test_normalize():
    assert normalize_text('  A  B  ') == 'A B'
    assert normalize_text('A\r\nB\rC') == 'A\nB\nC'


def test_format_duration():
    assert '秒' in format_duration(45.2)
    assert '分' in format_duration(125.5)


def test_cache_path():
    p = generate_cache_path('test.pdf')
    assert p.endswith('.json')


test('逗号分隔', test_read_comma)
test('换行分隔', test_read_newline)
test('混合分隔', test_read_mixed)
test('去重', test_read_dup)
test('分号分隔', test_read_semicolon)
test('中文逗号', test_read_chinese_comma)
test('中文分号', test_read_chinese_semicolon)
test('normalize_text', test_normalize)
test('format_duration', test_format_duration)
test('cache_path', test_cache_path)

# ==================== name_matcher.py ====================
print('\n[2/5] name_matcher.py')

from name_matcher import build_pattern, match_single_page, match_all_pages, get_match_statistics


def test_exact_match():
    pattern, name_set = build_pattern(['张三', '李四'])
    lower_map = {n.lower(): n for n in name_set}
    r = match_single_page('张三和李四', pattern, name_set, lower_map, 1)
    assert '张三' in r and '李四' in r


def test_no_substring_long():
    pattern, name_set = build_pattern(['张三', '张三丰'])
    lower_map = {n.lower(): n for n in name_set}
    r = match_single_page('张三丰来了', pattern, name_set, lower_map, 1)
    assert '张三丰' in r
    assert '张三' not in r


def test_no_substring_short():
    pattern, name_set = build_pattern(['张三', '张三丰'])
    lower_map = {n.lower(): n for n in name_set}
    r = match_single_page('张三来了', pattern, name_set, lower_map, 1)
    assert '张三' in r
    assert '张三丰' not in r


def test_multi_page():
    data = {1: '张三来了', 2: '李四来了', 3: '张三又来了'}
    r = match_all_pages(data, ['张三', '李四'])
    assert r['张三']['count'] == 2
    assert r['李四']['count'] == 1
    assert sorted(r['张三']['pages']) == [1, 3]


def test_no_match():
    data = {1: '王五来了'}
    r = match_all_pages(data, ['张三', '李四'])
    assert r['张三']['count'] == 0
    assert r['李四']['count'] == 0


def test_statistics():
    data = {1: '张三和李四', 2: '张三又来'}
    r = match_all_pages(data, ['张三', '李四', '王五'])
    stats = get_match_statistics(r)
    assert stats['total_names'] == 3
    assert stats['matched_count'] == 2
    assert stats['unmatched_count'] == 1
    assert stats['total_matches'] == 3


def test_case_insensitive():
    pattern, name_set = build_pattern(['John', 'Mary'], case_sensitive=False)
    lower_map = {n.lower(): n for n in name_set}
    r = match_single_page('john and MARY', pattern, name_set, lower_map, 1, case_sensitive=False)
    assert 'John' in r, f'应该返回原始大小写John，实际: {r}'
    assert 'Mary' in r, f'应该返回原始大小写Mary，实际: {r}'


def test_same_page_duplicates():
    pattern, name_set = build_pattern(['张三'])
    lower_map = {n.lower(): n for n in name_set}
    r = match_single_page('张三说张三很厉害', pattern, name_set, lower_map, 1)
    assert r['张三'] == [1]


def test_english_substring():
    pattern, name_set = build_pattern(['John', 'Johnson'])
    lower_map = {n.lower(): n for n in name_set}
    r = match_single_page('Johnson is here', pattern, name_set, lower_map, 1)
    assert 'Johnson' in r
    assert 'John' not in r


def test_english_exact():
    pattern, name_set = build_pattern(['John', 'Johnson'])
    lower_map = {n.lower(): n for n in name_set}
    r = match_single_page('John is here', pattern, name_set, lower_map, 1)
    assert 'John' in r
    assert 'Johnson' not in r


test('精确匹配', test_exact_match)
test('不匹配子串(长)', test_no_substring_long)
test('不匹配子串(短)', test_no_substring_short)
test('多页匹配', test_multi_page)
test('无匹配', test_no_match)
test('统计信息', test_statistics)
test('不区分大小写', test_case_insensitive)
test('同页去重', test_same_page_duplicates)
test('英文子串过滤', test_english_substring)
test('英文精确匹配', test_english_exact)

# ==================== pdf_parser.py ====================
print('\n[3/5] pdf_parser.py')

from pdf_parser import parse_pdf, get_pdf_info, clear_cache
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter


def test_pdf_parse():
    pdf_path = os.path.join(test_dir, 'test.pdf')
    c = canvas.Canvas(pdf_path, pagesize=letter)
    c.drawString(100, 700, 'Hello World')
    c.showPage()
    c.drawString(100, 700, 'Page Two')
    c.showPage()
    c.save()

    data = parse_pdf(pdf_path, use_cache=False)
    assert len(data) == 2
    assert 'Hello' in data[1]
    assert 'Page Two' in data[2]


def test_pdf_info():
    pdf_path = os.path.join(test_dir, 'test.pdf')
    info = get_pdf_info(pdf_path)
    assert info['page_count'] == 2


def test_pdf_cache():
    pdf_path = os.path.join(test_dir, 'test.pdf')
    parse_pdf(pdf_path, use_cache=True)
    data = parse_pdf(pdf_path, use_cache=True)
    assert len(data) == 2


def test_clear_cache():
    count = clear_cache(test_dir)
    assert count >= 0


test('PDF解析', test_pdf_parse)
test('PDF信息', test_pdf_info)
test('PDF缓存', test_pdf_cache)
test('清除缓存', test_clear_cache)

# ==================== excel_exporter.py ====================
print('\n[4/5] excel_exporter.py')

from excel_exporter import export_to_excel, get_output_filename
from openpyxl import load_workbook


def test_excel_export():
    result = {
        'A': {'count': 10, 'pages': [1, 2, 3]},
        'B': {'count': 5, 'pages': [4, 5]},
        'C': {'count': 0, 'pages': []}
    }
    stats = {
        'total_names': 3, 'matched_count': 2, 'unmatched_count': 1,
        'total_matches': 15, 'matched_names': ['A', 'B'], 'unmatched_names': ['C']
    }
    out = os.path.join(test_dir, 'out.xlsx')
    export_to_excel(result, out, stats=stats, duration=1.5)
    assert os.path.exists(out)


def test_excel_sort():
    result = {
        'C': {'count': 1, 'pages': [1]},
        'A': {'count': 10, 'pages': [1]},
        'B': {'count': 5, 'pages': [1]}
    }
    out = os.path.join(test_dir, 'sorted.xlsx')
    export_to_excel(result, out)
    wb = load_workbook(out)
    ws = wb['匹配结果']
    first_name = ws.cell(row=2, column=1).value
    assert first_name == 'A', f'应该是A，实际是{first_name}'


def test_excel_sheets():
    result = {'X': {'count': 1, 'pages': [1]}}
    stats = {
        'total_names': 1, 'matched_count': 1, 'unmatched_count': 0,
        'total_matches': 1, 'matched_names': ['X'], 'unmatched_names': []
    }
    out = os.path.join(test_dir, 'sheets.xlsx')
    export_to_excel(result, out, stats=stats)
    wb = load_workbook(out)
    assert '匹配结果' in wb.sheetnames
    assert '统计信息' in wb.sheetnames


def test_output_filename():
    name = get_output_filename(r'D:\test\doc.pdf')
    assert 'doc' in name
    assert '.xlsx' in name


test('Excel导出', test_excel_export)
test('Excel排序', test_excel_sort)
test('Excel工作表', test_excel_sheets)
test('输出文件名', test_output_filename)

# ==================== 集成测试 ====================
print('\n[5/5] 集成测试')


def test_full_pipeline():
    from utils import read_name_list
    from pdf_parser import parse_pdf
    from name_matcher import match_all_pages, get_match_statistics
    from excel_exporter import export_to_excel

    pdf_path = os.path.join(test_dir, 'full.pdf')
    c = canvas.Canvas(pdf_path, pagesize=letter)
    c.drawString(100, 700, 'John and Mary went home')
    c.showPage()
    c.drawString(100, 700, 'Bob saw John again')
    c.showPage()
    c.drawString(100, 700, 'Nobody here')
    c.showPage()
    c.save()

    name_file = os.path.join(test_dir, 'names.txt')
    with open(name_file, 'w', encoding='utf-8') as fp:
        fp.write('John,Mary,Bob')

    names = read_name_list(name_file)
    data = parse_pdf(pdf_path, use_cache=False)
    result = match_all_pages(data, names)
    stats = get_match_statistics(result)
    out = os.path.join(test_dir, 'final.xlsx')
    export_to_excel(result, out, stats=stats, duration=2.0)

    assert result['John']['count'] == 2
    assert result['Mary']['count'] == 1
    assert result['Bob']['count'] == 1
    assert os.path.exists(out)


def test_large_name_list():
    """测试大量人名"""
    from name_matcher import match_all_pages

    names = [f'Person{i}' for i in range(100)]
    data = {1: 'Person50 and Person99', 2: 'No one here'}
    r = match_all_pages(data, names)
    assert r['Person50']['count'] == 1
    assert r['Person99']['count'] == 1
    assert r['Person0']['count'] == 0


def test_empty_pdf_text():
    """测试空文本页"""
    from name_matcher import match_all_pages

    data = {1: '', 2: '   ', 3: 'John here'}
    r = match_all_pages(data, ['John'])
    assert r['John']['count'] == 1


test('完整流程', test_full_pipeline)
test('大量人名', test_large_name_list)
test('空文本页', test_empty_pdf_text)

# ==================== 清理 ====================
print('\n清理临时文件...')
shutil.rmtree(test_dir, ignore_errors=True)

# ==================== 结果 ====================
print('\n' + '=' * 60)
if errors:
    print(f'基准测试完成: {len(errors)} 个错误')
    for name, err in errors:
        print(f'  ✗ {name}: {err}')
    sys.exit(1)
else:
    print('基准测试完成: 全部通过 ✓')
print('=' * 60)
